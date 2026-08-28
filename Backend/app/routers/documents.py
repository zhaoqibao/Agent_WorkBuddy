"""文档上传 / 列表 / 预览（MinIO 存储 + 文本解析）。"""
from __future__ import annotations

import io
import uuid

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from sqlalchemy import select

from app.core.response import ok
from app.deps import CurrentUser, DBDep
from app.models import Document, KnowledgeDoc
from app.schemas import DocumentOut
from app.services.storage import get_presigned_url, put_object

router = APIRouter(prefix="/api/documents", tags=["documents"])

ALLOWED = {"docx", "xlsx", "pdf", "txt", "md", "csv"}
MAX_BYTES = 20 * 1024 * 1024  # 20 MB


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    try:
        if ext == "txt" or ext == "md" or ext == "csv":
            return data.decode("utf-8", errors="ignore")
        if ext == "docx":
            from docx import Document as Docx
            doc = Docx(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs)
        if ext == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.append("\t".join("" if c is None else str(c) for c in row))
            return "\n".join(parts)
        if ext == "pdf":
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                return "\n".join((p.extract_text() or "") for p in pdf.pages)
    except Exception:
        return ""
    return ""


@router.post("/upload")
async def upload(
    file: UploadFile = File(...),
    workspace_id: int = Form(...),
    knowledge_doc_id: int | None = Form(None),
    current: CurrentUser = None,
    db: DBDep = None,
):
    ext = file.filename.rsplit(".", 1)[-1].lower()
    if ext not in ALLOWED:
        raise HTTPException(status_code=400, detail=f"不支持的文件类型: {ext}")

    data = await file.read()
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=400, detail="文件超过 20MB 限制")

    # 若无知识库条目则自动创建
    if knowledge_doc_id is None:
        kd = KnowledgeDoc(user_id=current.id, workspace_id=workspace_id,
                          title=file.filename)
        db.add(kd)
        await db.flush()
        knowledge_doc_id = kd.id

    key = f"ws-{workspace_id}/u-{current.id}/{uuid.uuid4().hex}.{ext}"
    put_object(key, data, len(data), file.content_type or "application/octet-stream")

    text = extract_text(file.filename, data)
    doc = Document(
        knowledge_doc_id=knowledge_doc_id,
        user_id=current.id,
        workspace_id=workspace_id,
        original_name=file.filename,
        stored_path=key,
        file_type=ext,
        file_size=len(data),
        text_content=text,
        parse_status=1 if text else 2,
    )
    db.add(doc)
    await db.commit()
    await db.refresh(doc)
    return ok(DocumentOut.model_validate(doc).model_dump())


@router.get("")
async def list_documents(current: CurrentUser, db: DBDep, knowledge_doc_id: int | None = None):
    stmt = select(Document).where(
        Document.user_id == current.id, Document.deleted_at.is_(None)
    )
    if knowledge_doc_id is not None:
        stmt = stmt.where(Document.knowledge_doc_id == knowledge_doc_id)
    rows = (await db.scalars(stmt)).all()
    return ok([DocumentOut.model_validate(r).model_dump() for r in rows])


@router.get("/{doc_id}/preview")
async def preview(doc_id: int, current: CurrentUser, db: DBDep):
    doc = await db.get(Document, doc_id)
    if not doc or doc.deleted_at is not None or doc.user_id != current.id:
        raise HTTPException(status_code=404, detail="文档不存在")
    url = get_presigned_url(doc.stored_path, expires_minutes=15)
    return ok({"url": url, "original_name": doc.original_name, "text": doc.text_content})
